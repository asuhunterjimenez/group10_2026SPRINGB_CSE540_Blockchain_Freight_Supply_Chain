# apps/Payments/utils.py
import os
from decimal import Decimal
from django.conf import settings
from django.core.mail import EmailMessage
from docx import Document
from docx2pdf import convert

from apps.Payments.models import Payment

def generate_and_send_payment_receipt(payment_id):
    """
    Generates a payment receipt Word and PDF, and emails it to the customer.
    """
    try:
        payment = Payment.objects.get(id=payment_id)
    except Payment.DoesNotExist:
        print(f"Payment with id={payment_id} does not exist.")
        return False

    # Paths
    media_folder = os.path.join(settings.MEDIA_ROOT, "system_documents")
    os.makedirs(media_folder, exist_ok=True)

    word_template_path = os.path.join(media_folder, "payment_receipt.docx")
    word_filename = f"payment_receipt_{payment.transaction_id}.docx"
    word_path = os.path.join(media_folder, word_filename)

    pdf_filename = f"payment_receipt_{payment.transaction_id}.pdf"
    pdf_path = os.path.join(media_folder, pdf_filename)

    # Load Word template
    doc = Document(word_template_path)

    # Placeholder mapping
    placeholders = {
        "{{transaction_id}}": str(payment.transaction_id),
        "{{amount}}": str(payment.amount),
        "{{currency}}": str(payment.currency),
        "{{fx_fee}}": str(payment.fx_fee or Decimal("0.00")),
        "{{stripe_fee}}": str(payment.stripe_fee or Decimal("0.00")),
        "{{net_amount}}": str(payment.net_amount or Decimal("0.00")),
        "{{settled_amount}}": str(payment.settled_amount or Decimal("0.00")),
        "{{status}}": payment.status,
        "{{quote_request_id}}": str(payment.quote_request_id or "-"),
    }

    # Replace placeholders in paragraphs
    for p in doc.paragraphs:
        for key, val in placeholders.items():
            if key in p.text:
                inline = p.runs
                for i in range(len(inline)):
                    if key in inline[i].text:
                        inline[i].text = inline[i].text.replace(key, val)

    # Save Word file
    doc.save(word_path)

    # Convert Word to PDF
    convert(word_path, pdf_path)
    email_id=payment.user.email
    # Send email with PDF attachment
    email = EmailMessage(
        subject="Your Payment Receipt",
        body="Please find attached your payment receipt.",
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=[email_id],
    )
    email.attach_file(pdf_path)
    email.send(fail_silently=False)

    print(f"Payment receipt sent to {email_id}")
    return True
